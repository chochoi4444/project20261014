"""가상 장부로 끝까지: 읽기 → 매칭 → 총자산 → 저장 (CI에서도 실제 데이터 없이 돈다)"""
import json
import zipfile

import openpyxl

from make_fixture import make
from nav_total_assets.pipeline import run, save_workbook


def test_synthetic_end_to_end(tmp_path):
    d = make(tmp_path / "fx")
    expected = json.loads((d / "expected.json").read_text(encoding="utf-8"))
    target, ledgers, results, base = run(d / "대상.xlsx", [d / "장부"])
    assert base == "2026-09-30"
    assert all(not l.error for l in ledgers)
    by_name = {r.name: r for r in results}
    for name, value in expected.items():
        r = by_name[name]
        assert r.status == "auto" and r.include, (name, r.status, r.notes)
        assert round(r.ledger.value) == value
        assert r.ledger.method.endswith("당일잔액")
    # 통합 파일 하나에서 펀드 2개를 시트별로 읽는다
    for name, value, sheet in [("보람에너지풍력사모특별자산투자신탁", 3_000_000_000, "총계정원장_풍력"),
                               ("보람에너지태양광사모특별자산투자신탁", 1_500_000_000, "총계정원장_태양광")]:
        r = by_name[name]
        assert r.ledger and r.ledger.sheet == sheet and round(r.ledger.value) == value, (name, r.status, r.notes)
        assert r.status in ("auto", "review")
    assert by_name["예시법인용MMF 1(국공채)C"].status == "skip"
    assert by_name["가온대체투자전문투자형사모부동산투자신탁102호"].status == "nomatch"  # 101호 장부에 붙으면 안 된다
    out = tmp_path / "out.xlsx"
    assert save_workbook(target, results, out) == sum(r.include for r in results) >= 2
    ws = openpyxl.load_workbook(out).worksheets[0]
    assert ws["D2"].value == 5_100_000_000 and ws["D3"].value == 21_000_000_000
    assert ws["D4"].value is None and ws["D5"].value == "제외"
    with zipfile.ZipFile(d / "대상.xlsx") as a, zipfile.ZipFile(out) as b:
        assert a.namelist() == b.namelist()
        changed = {n for n in a.namelist() if a.read(n) != b.read(n)}
        # workbook.xml은 '열 때 다시 계산' 설정이 이미 있으면 바뀌지 않는다
        assert "xl/worksheets/sheet1.xml" in changed and changed <= {"xl/workbook.xml", "xl/worksheets/sheet1.xml"}
