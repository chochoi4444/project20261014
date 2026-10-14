"""기준가격대장(신탁재산명세부) 읽기와 총자산 산출.

모든 형식(xlsx · xls · csv · html로 된 xls · 텍스트 PDF)을 '행 = 셀 목록 [Cell(c, v, r)]'로 통일한 뒤
같은 규칙으로 계정을 찾는다. 웹 버전과 같은 규칙이다.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from .engine import nfkc, norm_label, norm_name


@dataclass
class Cell:
    c: float            # 열 위치 (엑셀: 열 번호, PDF: 글자 가운데 x좌표)
    v: object           # 숫자(float) 또는 문자열
    r: float | None = None  # PDF 전용: 글자 오른쪽 끝 x좌표


# ---------- 숫자·날짜 ----------
_NUM = re.compile(r"^-?[\d,]+(\.\d+)?$")
_NUM_GROUPED = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$")


def to_num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if v == v and v not in (float("inf"), float("-inf")) else None
    if not isinstance(v, str):
        return None
    t = re.sub(r"\s", "", v.strip())
    if not re.search(r"\d", t):
        return None
    neg = False
    if re.fullmatch(r"\(.*\)", t):
        neg, t = True, t[1:-1]
    if not _NUM.match(t):
        return None
    if "," in t and not _NUM_GROUPED.match(t):
        return None
    try:
        n = float(t.replace(",", ""))
    except ValueError:
        return None
    return -n if neg else n


def _iso(y, m, d):
    return f"{y:04d}-{m:02d}-{d:02d}"


def _valid(y, m, d):
    return 2000 <= y <= 2099 and 1 <= m <= 12 and 1 <= d <= 31


def date_from_text(s) -> str | None:
    t = nfkc(s)
    found = None
    for m in re.finditer(r"(20\d{2})\s*[.\-/년]\s*(\d{1,2})\s*[.\-/월]\s*(\d{1,2})", t):
        y, mo, d = int(m[1]), int(m[2]), int(m[3])
        if _valid(y, mo, d):
            found = _iso(y, mo, d)
    if found:
        return found
    for m in re.finditer(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)", t):
        y, mo, d = int(m[1]), int(m[2]), int(m[3])
        if _valid(y, mo, d):
            found = _iso(y, mo, d)
    return found


def date_from_filename(s) -> str | None:
    d = date_from_text(s)
    if d:
        return d
    m = re.search(r"(?<!\d)(\d{2})(\d{2})(\d{2})(?!\d)", nfkc(s))
    if m and 15 <= int(m[1]) <= 45 and _valid(2000 + int(m[1]), int(m[2]), int(m[3])):
        return _iso(2000 + int(m[1]), int(m[2]), int(m[3]))
    return None


def serial_to_iso(n) -> str | None:
    if isinstance(n, bool) or not isinstance(n, (int, float)) or n < 30000 or n > 80000:
        return None
    d = dt.datetime(1970, 1, 1) + dt.timedelta(milliseconds=round((n - 25569) * 86400000))
    return _iso(d.year, d.month, d.day)


def clean_file_name(name: str) -> str:
    t = re.sub(r"\.[^.]+$", "", nfkc(name))
    t = re.sub(r"(\d{6,8})[\s_\-]*(\(\d+\)|_\d)\s*$", r"\1", t)
    t = re.sub(r"\s*\(\d+\)\s*$", "", t)
    t = re.sub(r"20\d{6}|(?<!\d)\d{6}(?!\d)", " ", t)
    t = re.sub(r"(20\d{2}\s*년?\s*)?[1-4]\s*(Q|분기)", " ", t, flags=re.I)
    t = re.sub(r"20\d{2}\s*년", " ", t)
    t = re.sub(r"기준가격대장|기준가대장|신탁재산명세부|재산명세부|명세부|계정원장|원장|조회|사본|복사본", " ", t)
    return re.sub(r"\s+", " ", re.sub(r"_+", " ", t)).strip()


# ---------- 형식별 읽기 ----------

def _cell(j, v):
    if v is None or v == "":
        return None
    if isinstance(v, str):
        n = to_num(v)
        return Cell(j, v if n is None else n)
    if isinstance(v, bool):
        return Cell(j, str(v))
    if isinstance(v, (int, float)):
        return Cell(j, float(v))
    if isinstance(v, (dt.datetime, dt.date)):
        # 엑셀 날짜 셀은 일련번호로 바꿔 다른 형식과 똑같이 다룬다
        base = dt.datetime(1899, 12, 30)
        d = v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day)
        return Cell(j, (d - base).total_seconds() / 86400)
    return Cell(j, str(v))


def _rows_from_values(value_rows) -> list[list[Cell]]:
    return [[c for c in (_cell(j, v) for j, v in enumerate(r or [])) if c is not None] for r in value_rows]


def sheets_from_xlsx(path: Path) -> list[list[list[Cell]]]:
    import openpyxl
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [_rows_from_values(ws.iter_rows(values_only=True)) for ws in wb.worksheets]
    finally:
        wb.close()


def sheets_from_xls(path: Path) -> list[list[list[Cell]]]:
    import xlrd
    book = xlrd.open_workbook(str(path), logfile=io.StringIO())
    out = []
    for sh in book.sheets():
        rows = []
        for i in range(sh.nrows):
            vals = []
            for j in range(sh.ncols):
                c = sh.cell(i, j)
                if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    vals.append(None)
                elif c.ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_DATE):
                    vals.append(float(c.value))
                elif c.ctype == xlrd.XL_CELL_BOOLEAN:
                    vals.append(bool(c.value))
                elif c.ctype == xlrd.XL_CELL_ERROR:
                    vals.append(None)
                else:
                    vals.append(c.value)
            rows.append(vals)
        out.append(_rows_from_values(rows))
    return out


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def sheets_from_csv(path: Path):
    text = _decode(path.read_bytes())
    return [_rows_from_values(list(csv.reader(io.StringIO(text))))]


class _TableParser(HTMLParser):
    """사내 시스템이 .xls 확장자로 내보내는 HTML 표를 읽는다"""

    def __init__(self):
        super().__init__()
        self.tables, self._row, self._cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr":
            self._row = []
        elif tag in ("td", "th"):
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(re.sub(r"\s+", " ", "".join(self._cell)).strip() or None)
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if not self.tables:
                self.tables.append([])
            self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def sheets_from_html(path: Path):
    p = _TableParser()
    p.feed(_decode(path.read_bytes()))
    return [_rows_from_values(t) for t in p.tables if t]


def rows_from_pdf(path: Path) -> list[list[Cell]]:
    import pdfplumber
    rows, text_len = [], 0
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            # 칸 안의 띄어쓰기는 한 덩어리로 두고(keep_blank_chars), 칸 사이의 넓은 간격에서만 나눈다
            words = page.extract_words(keep_blank_chars=True, x_tolerance=1.5, y_tolerance=2.5, use_text_flow=False)
            words = [w for w in words if w["text"].strip()]
            text_len += sum(len(w["text"].strip()) for w in words)
            lines: list[dict] = []
            for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
                ln = next((l for l in lines if abs(l["y"] - w["top"]) <= 2.5), None)
                if ln:
                    ln["items"].append(w)
                else:
                    lines.append({"y": w["top"], "items": [w]})
            for ln in sorted(lines, key=lambda l: l["y"]):
                row = []
                for w in sorted(ln["items"], key=lambda w: w["x0"]):
                    s = w["text"].strip()
                    n = to_num(s)
                    row.append(Cell((w["x0"] + w["x1"]) / 2, s if n is None else n, w["x1"]))
                rows.append(row)
    if text_len < 20:
        raise ValueError("텍스트가 없는 PDF(스캔본)라 읽을 수 없습니다. 수기로 확인하세요.")
    return rows


# ---------- 계정 찾기 ----------
FIELDS = {
    "total": ["신탁재산총자산", "신탁자산총자산", "신탁재산총자산(감사보수기준)", "신탁자산총자산(감사보수기준)", "총자산", "총자산액", "자산총액", "펀드총자산"],
    "net": ["신탁재산순자산총액", "신탁자산순자산총액", "순자산총액", "신탁재산순자산", "순자산"],
    "liab": ["신탁자산총부채", "신탁재산총부채", "부채", "부채합계", "부채총계", "부채총액", "총부채"],
    "asset": ["자산", "자산합계", "자산총계"],
    "nav": ["수익증권기준가격", "수익증권기준가액", "기준가격", "기준가"],
    "units": ["수익증권잔존좌수", "잔존좌수", "수익증권좌수", "발행좌수", "총좌수"],
}
AMOUNT_FIELDS = ("total", "net", "liab", "asset")


def _is_str(c: Cell):
    return isinstance(c.v, str)


def _is_num(c: Cell):
    return isinstance(c.v, float)


def header_at(row: list[Cell]):
    """'당일잔액'이 있는 차변/대변 머리행"""
    texts = [c for c in row if _is_str(c)]
    joined = "".join(norm_label(c.v) for c in texts)
    if not (re.search("당일", joined) and re.search("잔액|잔고", joined)):
        return None
    dang = [c.c for c in texts if "당일" in norm_label(c.v) and "증감" not in norm_label(c.v)]
    if not dang:
        return None
    # PDF 셀은 오른쪽 끝(r)을 가진다 → 머리글 시작 위치 = 2·중심 − 오른쪽 끝
    starts = sorted(({"c": c.c, "start": 2 * c.c - c.r} for c in row if c.r is not None), key=lambda x: x["start"])
    return {"cols": [c.c for c in row], "left": min(dang), "right": max(dang), "starts": starts}


def extract_from_rows(rows: list[list[Cell]]):
    out = {"fields": {}, "header_row": None, "has_header": False}
    header = None
    found: dict[str, dict] = {}
    for i, row in enumerate(rows):
        texts = [c for c in row if _is_str(c)]
        joined = "".join(norm_label(c.v) for c in texts)
        h = header_at(row)
        if h:
            header = h
            out["has_header"] = True
            if out["header_row"] is None:
                out["header_row"] = i
            continue
        if not texts:
            continue
        labels = {norm_label(c.v) for c in texts}
        if len(texts) > 1:
            labels.add(joined)
        for fld, names in FIELDS.items():
            pri = next((k for k, n in enumerate(names) if n in labels), -1)
            if pri < 0:
                continue
            if fld in found and found[fld]["pri"] <= pri:
                continue
            left = right = first = None
            label_col = max(c.c for c in texts if norm_label(c.v) in names or joined in labels)
            for cell in row:
                if not _is_num(cell):
                    continue
                if header:
                    near = header["cols"][0]
                    if cell.r is not None and header["starts"]:
                        # PDF: 숫자는 오른쪽 정렬 — 숫자 오른쪽 끝보다 왼쪽에서 시작하는 마지막 머리글이 그 열
                        hs = [x for x in header["starts"] if x["start"] <= cell.r + 2]
                        near = hs[-1]["c"] if hs else header["starts"][0]["c"]
                    else:
                        d = float("inf")
                        for hc in header["cols"]:
                            if abs(hc - cell.c) < d:
                                d, near = abs(hc - cell.c), hc
                    if near == header["left"] and left is None:
                        left = cell.v
                    elif near == header["right"] and right is None:
                        right = cell.v
                elif cell.c > label_col and first is None:
                    first = cell.v
            if header:
                pref = [right, left] if fld == "liab" else [left, right]
                v = next((x for x in pref if x is not None and x != 0), None)
                if v is None:
                    v = next((x for x in pref if x is not None), None)
            else:
                v = first
            if v is None:
                continue
            found[fld] = {"pri": pri, "v": v, "label": names[pri]}
    out["fields"] = found
    return out


def _fmt_code(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def extract_meta(rows: list[list[Cell]], upto: int | None):
    meta: dict = {}
    lim = min(len(rows), 15 if upto is None else upto)
    for i in range(lim):
        row = rows[i]
        for k, cell in enumerate(row):
            if not _is_str(cell):
                continue
            raw = nfkc(cell.v).strip()
            tl = re.sub(r"\s", "", raw)
            if tl.startswith("펀드명"):
                rest = re.sub(r"^\s*펀\s*드\s*명\s*[:：▶]?\s*", "", raw).strip()
                if re.search(r"[가-힣A-Za-z]", rest):
                    meta.setdefault("name", rest)
                else:
                    for c in row[k + 1:]:
                        if _is_num(c):
                            meta.setdefault("code", _fmt_code(c.v))
                        elif re.search(r"[가-힣A-Za-z]", c.v):
                            meta.setdefault("name", nfkc(c.v).strip())
                            break
                continue
            m = re.fullmatch(r"펀드(코드)?[:：]([A-Za-z0-9\-]+)", tl)
            if m:
                meta.setdefault("code", m[2])
                continue
            if re.search(r"(기준가?산출일자|기준일자|기준일|작업일자|기준가격일자?|조회일자?|평가일자?)", tl):
                d = date_from_text(raw)
                if not d:
                    for c in row[k + 1:]:
                        d = serial_to_iso(c.v) if _is_num(c) else date_from_text(c.v)
                        if d:
                            break
                if d:
                    meta.setdefault("date", d)
            if "manager" not in meta and re.search(r"(자산운용|운용㈜|투자운용|ASSETMANAGEMENT|INVESTMENT|INVESTMENTTRUST)", tl, re.I):
                meta["manager"] = raw
            m = re.search(r"단위[:：]?\(?(천원|백만원|억원|원)", tl)
            if m:
                meta["unit"] = m[1]
    if "date" not in meta:
        for i in range(lim):
            d = next((date_from_text(c.v) for c in rows[i] if _is_str(c) and date_from_text(c.v)), None)
            if d:
                meta["date"] = d
                break
    return meta


UNIT_MULT = {"원": 1, "천원": 1e3, "백만원": 1e6, "억원": 1e8}


@dataclass
class Ledger:
    id: int
    path: Path
    file_name: str
    name: str | None = None
    file_title: str = ""
    name_cands: list = field(default_factory=list)
    manager: str = ""
    code: str = ""
    date: str | None = None
    key: str = ""
    fields: dict = field(default_factory=dict)
    meta: dict = field(default_factory=dict)
    error: str | None = None
    value: float | None = None
    method: str = ""
    check: tuple = ("mute", "")

    def as_match_input(self) -> dict:
        return {"id": self.id, "name_cands": self.name_cands, "manager": self.manager,
                "name": self.name, "file_title": self.file_title}


def _sheets(path: Path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        return [rows_from_pdf(path)]
    head = path.read_bytes()[:512].lstrip()
    if head[:1] == b"<":
        return sheets_from_html(path)
    if ext == ".csv":
        return sheets_from_csv(path)
    if head[:4] == b"PK\x03\x04":
        return sheets_from_xlsx(path)
    if head[:8] == bytes.fromhex("D0CF11E0A1B11AE1"):
        return sheets_from_xls(path)
    raise ValueError("엑셀로 열 수 없습니다. 암호·DRM이 걸려 있는지 확인하세요.")


def parse_ledger(path: Path, lid: int) -> Ledger:
    path = Path(path)
    L = Ledger(id=lid, path=path, file_name=path.name)
    try:
        try:
            sheets = _sheets(path)
        except ValueError:
            raise
        except Exception as e:  # 손상·암호·DRM 파일
            raise ValueError(f"파일을 열 수 없습니다 ({type(e).__name__}). 암호·DRM이 걸려 있는지 확인하세요.")
        best = None
        for rows in sheets:
            ex = extract_from_rows(rows)
            meta = extract_meta(rows, ex["header_row"])
            rank = len(ex["fields"]) + (2 if ex["has_header"] else 0)
            if best is None or rank > best[0]:
                best = (rank, ex, meta)
        if best is None:
            raise ValueError("내용이 비어 있습니다.")
        _, ex, meta = best
        L.meta = meta
        mult = UNIT_MULT.get(meta.get("unit"), 1)
        for k, f in ex["fields"].items():
            L.fields[k] = {**f, "v": f["v"] * mult if k in AMOUNT_FIELDS else f["v"]}
    except ValueError as e:
        L.error = str(e)
    L.name = L.meta.get("name")
    L.file_title = clean_file_name(path.name)
    L.name_cands = [x for x in (L.name, L.file_title) if x]
    L.manager = L.meta.get("manager", "")
    L.code = L.meta.get("code", "")
    L.date = L.meta.get("date") or date_from_filename(path.name)
    L.key = norm_name(L.name or L.file_title)
    compute_total(L)
    return L


def _fmt(n):
    return f"{round(n):,}"


def compute_total(L: Ledger):
    """총자산 = 신탁재산총자산 계정 → 순자산총액+부채 → 자산합계 순, 순자산+부채로 교차검증"""
    if L.error:
        L.value, L.check = None, ("bad", "읽기 실패")
        return
    f = lambda k: (L.fields.get(k) or {}).get("v")
    total, net, liab, asset = f("total"), f("net"), f("liab"), f("asset")
    near = lambda a, b: abs(a - b) <= 10
    if total is not None:
        L.value, L.method = total, L.fields["total"]["label"] + " 계정"
        if net is not None and liab is not None:
            L.check = ("ok", "순자산+부채 일치") if near(total, net + liab) else ("warn", f"순자산+부채와 {_fmt(total - net - liab)}원 차이")
        else:
            L.check = ("mute", "교차검증 값 없음")
    elif net is not None and liab is not None:
        L.value = net + liab
        L.method = f"{L.fields['net']['label']} + {L.fields['liab']['label']}"
        if asset is not None:
            L.check = ("ok", "자산합계와 일치") if near(asset, L.value) else ("warn", "자산합계와 다름(장부가 확인)")
        else:
            L.check = ("mute", "교차검증 값 없음")
    elif asset is not None:
        L.value, L.method = asset, L.fields["asset"]["label"] + " (장부가일 수 있음)"
        L.check = ("warn", "평가 반영 여부 확인")
    else:
        L.value, L.method, L.check = None, "", ("bad", "총자산 계정을 찾지 못함")
