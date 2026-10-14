"""대상 엑셀 읽기와 저장.

저장은 원본 xlsx(zip)를 그대로 복사하면서 총자산 칸의 XML만 바꾼다.
서식·데이터 유효성·다른 시트·인쇄 설정은 손대지 않는다 (openpyxl로 다시 저장하면 일부가 사라진다).
"""
from __future__ import annotations

import datetime as dt
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from .ledger import date_from_text, serial_to_iso, to_num


@dataclass
class TargetRow:
    excel_row: int
    name: str
    code: str = ""
    cls: str = ""
    existing: object = None


@dataclass
class Target:
    path: Path
    sheet_name: str
    sheet_index: int
    asset_col_letter: str
    asset_header: str
    base_date: str | None
    rows: list = field(default_factory=list)


def _date_value(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.year:04d}-{v.month:02d}-{v.day:02d}"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return serial_to_iso(v)
    return date_from_text(v) if v is not None else None


def parse_target(path: Path) -> Target:
    import openpyxl
    from openpyxl.utils import get_column_letter
    import warnings
    path = Path(path)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    for si, ws in enumerate(wb.worksheets):
        aoa = [list(r) for r in ws.iter_rows(min_row=1, min_col=1, values_only=True)]
        for hr in range(min(15, len(aoa))):
            row = aoa[hr]

            def find(rx, flags=0):
                for j, v in enumerate(row):
                    if isinstance(v, str) and re.search(rx, re.sub(r"\s", "", v), flags):
                        return j
                return -1

            name_col = find(r"^(종목명|펀드명|상품명)$")
            if name_col < 0:
                name_col = find(r"종목명|펀드명")
            asset_col = find(r"총자산")
            if name_col < 0 or asset_col < 0:
                continue
            code_col, cls_col, date_col = find(r"표준코드|ISIN", re.I), find(r"유가증권분류|분류"), find(r"^기준일")
            base_date = date_from_text(row[asset_col])
            rows = []
            for i in range(hr + 1, len(aoa)):
                r = aoa[i]
                name = r[name_col] if name_col < len(r) else None
                if not isinstance(name, str) or not name.strip():
                    continue
                if not base_date and date_col >= 0:
                    base_date = _date_value(r[date_col])
                get = lambda c: (r[c] if 0 <= c < len(r) and r[c] is not None else "")
                rows.append(TargetRow(excel_row=i + 1, name=name.strip(), code=str(get(code_col)), cls=str(get(cls_col)),
                                      existing=r[asset_col] if asset_col < len(r) else None))
            wb.close()
            return Target(path=path, sheet_name=ws.title, sheet_index=si, asset_col_letter=get_column_letter(asset_col + 1),
                          asset_header=re.sub(r"\s+", " ", str(row[asset_col])), base_date=base_date, rows=rows)
    wb.close()
    raise ValueError("‘종목명’과 ‘총자산’ 머리글이 있는 시트를 찾지 못했습니다.")


def is_text_mark(v) -> bool:
    """총자산 칸에 '제외' 같은 글자가 적힌 행"""
    return isinstance(v, str) and bool(v.strip()) and to_num(v) is None


def has_number(v) -> bool:
    return (isinstance(v, (int, float)) and not isinstance(v, bool)) or (isinstance(v, str) and to_num(v) is not None)


# ---------- 저장: 원본 XML에서 총자산 칸만 교체 ----------

def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def set_cell_number(xml: str, ref: str, value: int) -> str:
    row_num = int(re.search(r"\d+$", ref).group())
    col = re.match(r"^[A-Z]+", ref).group()
    cell_xml = lambda attrs: f'<c r="{ref}"{attrs}><v>{value}</v></c>'
    rm = re.search(rf'<row\b[^>]*\br="{row_num}"[^>]*?(/>|>)', xml)
    if not rm:
        # 행이 없으면 순서에 맞춰 새로 만든다
        if "<sheetData/>" in xml:
            return xml.replace("<sheetData/>", f'<sheetData><row r="{row_num}">{cell_xml("")}</row></sheetData>')
        at = xml.find("</sheetData>")
        for m in re.finditer(r'<row\b[^>]*\br="(\d+)"', xml):
            if int(m.group(1)) > row_num:
                at = m.start()
                break
        return xml[:at] + f'<row r="{row_num}">{cell_xml("")}</row>' + xml[at:]
    if rm.group(1) == "/>":
        opened = rm.group(0)[:-2] + ">"
        return xml[:rm.start()] + opened + cell_xml("") + "</row>" + xml[rm.end():]
    start = rm.end()
    end = xml.find("</row>", start)
    inner = xml[start:end]
    cm = re.search(rf'<c\b([^>]*?)\br="{ref}"([^>]*?)(/>|>[\s\S]*?</c>)', inner)
    if cm:
        attrs = re.sub(r'\s+t="[^"]*"', "", cm.group(1) + cm.group(2)).rstrip()
        inner = inner[:cm.start()] + cell_xml(" " + attrs.strip() if attrs else "") + inner[cm.end():]
    else:
        want, at = _col_index(col), len(inner)
        for m in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"', inner):
            if _col_index(m.group(1)) > want:
                at = m.start()
                break
        inner = inner[:at] + cell_xml("") + inner[at:]
    return xml[:start] + inner + xml[end:]


def _sheet_part(zin: zipfile.ZipFile, sheet_index: int) -> str:
    wb_xml = zin.read("xl/workbook.xml").decode("utf-8")
    rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    sheets = re.findall(r"<sheet\b[^>]*>", wb_xml)
    rid = re.search(r'r:id="([^"]+)"', sheets[sheet_index]).group(1)
    rel = next(t for t in re.findall(r"<Relationship\b[^>]*>", rels) if f'Id="{rid}"' in t)
    target = re.search(r'Target="([^"]+)"', rel).group(1)
    return target[1:] if target.startswith("/") else "xl/" + re.sub(r"^\./", "", target)


def write_values(src: Path, dst: Path, sheet_index: int, values: dict[str, float]) -> int:
    """values: {"W31": 9473493391, ...} → 새 파일 dst에 저장, 바꾼 칸 수를 돌려준다"""
    src, dst = Path(src), Path(dst)
    if src.resolve() == dst.resolve():
        raise ValueError("원본 파일은 덮어쓰지 않습니다. 다른 이름으로 저장하세요.")
    with zipfile.ZipFile(src) as zin:
        part = _sheet_part(zin, sheet_index)
        xml = zin.read(part).decode("utf-8")
        for ref, v in values.items():
            xml = set_cell_number(xml, ref, int(v + 0.5) if v >= 0 else -int(-v + 0.5))
        wb_xml = zin.read("xl/workbook.xml").decode("utf-8")
        # 열 때 다시 계산하도록
        if re.search(r"<calcPr\b", wb_xml):
            wb_xml = re.sub(r"<calcPr\b([^>]*?)(/?)>",
                            lambda m: f'<calcPr{re.sub(r"\s*fullCalcOnLoad=\"[^\"]*\"", "", m.group(1))} fullCalcOnLoad="1"{m.group(2)}>',
                            wb_xml, count=1)
        else:
            wb_xml = wb_xml.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')
        with zipfile.ZipFile(dst, "w") as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename == part:
                    data = xml.encode("utf-8")
                elif info.filename == "xl/workbook.xml":
                    data = wb_xml.encode("utf-8")
                zout.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED)
    return len(values)
