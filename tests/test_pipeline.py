"""더미데이터 정답지와 결과를 비교하는 회귀 테스트."""
from pathlib import Path

from openpyxl import load_workbook

from aggregator import extract_report, fill_form

DATA = Path(__file__).resolve().parents[1] / "dummy_data"


def run():
    reports = [extract_report(p.read_bytes(), p.name) for p in sorted((DATA / "input").glob("*.xlsx"))]
    form = (DATA / "target" / "당사양식_수익증권평가_빈양식.xlsx").read_bytes()
    mapping = (DATA / "master" / "종목_운용사_매핑표.xlsx").read_bytes()
    return reports, fill_form(form, reports, mapping)


def test_every_report_recognized():
    reports, _ = run()
    assert reports and all(r.ok and r.format in "ABC" for r in reports), [(r.filename, r.warnings) for r in reports]


def test_matches_answer_key():
    import io

    _, (filled, results, unmatched) = run()
    got = load_workbook(io.BytesIO(filled)).worksheets[0]
    want = load_workbook(DATA / "answer" / "당사양식_정답.xlsx").worksheets[0]
    col = [c.value for c in want[1]].index(next(h for h in [c.value for c in want[1]] if "총자산" in h)) + 1
    for r in range(2, want.max_row + 1):
        assert got.cell(r, col).value == want.cell(r, col).value, want.cell(r, 4).value
    assert not unmatched
    assert {x.status for x in results} == {"완료", "제외", "미수령"}


def test_works_without_mapping_file():
    reports, _ = run()
    form = (DATA / "target" / "당사양식_수익증권평가_빈양식.xlsx").read_bytes()
    _, results, _ = fill_form(form, reports, None)
    assert sum(r.status == "완료" for r in results) == 8


def test_name_matching_normalizes_roman_numerals_and_case():
    from aggregator.extract import match_key

    assert match_key("가상 USDL III Levered Note 1호") == match_key("가상 USDL Ⅲ LEVERED NOTE 1호")
    assert match_key("(주)가상테크") == match_key("가상테크 주식회사")
