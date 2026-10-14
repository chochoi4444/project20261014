"""운용사 보고서(엑셀)를 읽어 양식을 판별하고 총자산을 추출한다.

지원 양식
  A: 좌우형 대장   - '작업일자 ▶' / 코드 100000 '자산합계' 행의 차변 당일잔액
  B: 한 줄 헤더 표 - '계정과목명' 헤더 / 코드 1000000000 '[[ 자산 ]]' 행의 당일차변잔액
  C: 요약형        - '구분/금액' 세로 표 / '총자산' 행, '(단위: 백만원)' 등 단위 환산
"""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

Grid = list[list[object]]

UNIT_MULTIPLIER = {"원": 1, "천원": 1_000, "백만원": 1_000_000, "억원": 100_000_000}


@dataclass
class ReportResult:
    filename: str
    format: str = "미인식"
    fund_code: str | None = None
    fund_name: str | None = None
    base_date: date | None = None
    total_asset: int | None = None
    liabilities: int | None = None
    net_asset: int | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.total_asset is not None


# ---------------------------------------------------------------- 파일 읽기
def load_grid(data: bytes, filename: str) -> Grid:
    """첫 번째 시트를 2차원 리스트로 읽는다 (.xlsx / .xls)."""
    if filename.lower().endswith(".xls"):
        import xlrd

        sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
        return [sheet.row_values(r) for r in range(sheet.nrows)]

    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).worksheets[0]
    return [list(row) for row in ws.iter_rows(values_only=True)]


# ---------------------------------------------------------------- 정규화 유틸
def norm(text: object) -> str:
    """비교용 문자열: 공백·괄호·특수문자 제거."""
    return re.sub(r"[\s\[\]()（）▶:·_\-]", "", str(text or ""))


def match_key(text: object) -> str:
    """종목명 매칭용 키: 유니코드 호환 정규화(Ⅲ→III), 대문자, 법인 표기 제거."""
    s = unicodedata.normalize("NFKC", str(text or "")).upper()
    s = re.sub(r"\(주\)|주식회사", "", s)
    return norm(s)


def to_number(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        n = float(s)
    except ValueError:
        return None
    return -n if neg else n


def parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value or "")
    m = re.search(r"(\d{4})\s*[년.\-/]\s*(\d{1,2})\s*[월.\-/]\s*(\d{1,2})", s)
    if m:
        return date(*map(int, m.groups()))
    m = re.search(r"(?<!\d)(\d{4})(\d{2})(\d{2})(?!\d)", s)
    if m:
        try:
            return date(*map(int, m.groups()))
        except ValueError:
            pass
    m = re.search(r"(?<!\d)(\d{2})(\d{2})(\d{2})(?!\d)", s)
    if m:
        y, mo, d = map(int, m.groups())
        try:
            return date(2000 + y, mo, d)
        except ValueError:
            pass
    return None


def cells(grid: Grid):
    for r, row in enumerate(grid):
        for c, v in enumerate(row):
            if v not in (None, ""):
                yield r, c, v


def find_cell(grid: Grid, keyword: str) -> tuple[int, int] | None:
    key = norm(keyword)
    for r, c, v in cells(grid):
        if key in norm(v):
            return r, c
    return None


def right_values(grid: Grid, r: int, c: int) -> list[object]:
    return [v for v in grid[r][c + 1:] if v not in (None, "")]


def find_row(grid: Grid, codes: tuple[str, ...], names: tuple[str, ...]) -> int | None:
    """계정코드 우선, 없으면 계정명 키워드로 행을 찾는다 (규칙 R2, R3)."""
    for r, row in enumerate(grid):
        for v in row:
            s = str(v).strip() if v is not None else ""
            if s.endswith(".0"):
                s = s[:-2]
            if s in codes:
                return r
    keys = [norm(n) for n in names]
    for r, row in enumerate(grid):
        for v in row:
            if isinstance(v, str) and norm(v) in keys:
                return r
    return None


def num_at(grid: Grid, r: int | None, c: int | None) -> int | None:
    if r is None or c is None or c >= len(grid[r]):
        return None
    n = to_number(grid[r][c])
    return None if n is None else int(round(n))


# ---------------------------------------------------------------- 양식 판별
def detect_format(grid: Grid) -> str:
    """시트 안 표식 단어로 양식 판별 (규칙 R1)."""
    text = " ".join(norm(v) for _, _, v in cells(grid[:60]))
    if "작업일자" in text and "계정명" in text:
        return "A"
    if "계정과목명" in text and "당일차변잔액" in text:
        return "B"
    if "총자산" in text or "AUM" in text.upper():
        return "C"
    return "미인식"


# ---------------------------------------------------------------- 양식별 추출
def _extract_a(grid: Grid, res: ReportResult) -> None:
    if pos := find_cell(grid, "작업일자"):
        vals = right_values(grid, *pos)
        res.base_date = parse_date(vals[0]) if vals else None
    if pos := find_cell(grid, "펀드명"):
        vals = right_values(grid, *pos)
        if vals:
            code = to_number(vals[0])
            res.fund_code = str(int(code)) if code is not None else str(vals[0])
            res.fund_name = str(vals[1]) if len(vals) > 1 else None
    # '당일잔액' 헤더: 첫 번째 = 차변, 마지막 = 대변
    hdr = next((r for r, row in enumerate(grid) if sum(norm(v) == "당일잔액" for v in row) >= 1), None)
    if hdr is None:
        res.warnings.append("'당일잔액' 헤더를 찾지 못함")
        return
    cols = [c for c, v in enumerate(grid[hdr]) if norm(v) == "당일잔액"]
    debit, credit = cols[0], cols[-1]
    res.total_asset = num_at(grid, find_row(grid, ("100000",), ("자산합계",)), debit)
    res.liabilities = num_at(grid, find_row(grid, ("200000",), ("부채합계",)), credit)
    res.net_asset = num_at(grid, find_row(grid, ("700001",), ("순자산총액",)), credit)


def _extract_b(grid: Grid, res: ReportResult) -> None:
    hdr = next(r for r, row in enumerate(grid) if any(norm(v) == "당일차변잔액" for v in row))
    col = {norm(v): c for c, v in enumerate(grid[hdr]) if v}
    debit, credit = col.get("당일차변잔액"), col.get("당일대변잔액")
    res.total_asset = num_at(grid, find_row(grid, ("1000000000",), ("자산", "자산총계")), debit)
    res.liabilities = num_at(grid, find_row(grid, ("2000000000",), ("부채", "부채총계")), credit)
    res.net_asset = num_at(grid, find_row(grid, (), ("순자산", "순자산총액")), credit)
    # 펀드명·기준일은 파일명에서 (규칙 R6): 운용사_펀드명_YYMMDD.xlsx
    stem = Path(res.filename).stem
    parts = stem.split("_")
    res.base_date = parse_date(stem)
    if len(parts) >= 2:
        res.fund_name = parts[1]


def _extract_c(grid: Grid, res: ReportResult) -> None:
    unit = 1
    for _, _, v in cells(grid[:10]):
        s = str(v)
        if m := re.search(r"단위\s*[:：]?\s*(백만원|천원|억원|원)", s):
            unit = UNIT_MULTIPLIER[m.group(1)]
        if res.base_date is None and "기준일" in s:
            res.base_date = parse_date(s)
        if m := re.search(r"펀드코드\s*[:：]\s*(\S+)", s):
            res.fund_code = m.group(1)
    title = grid[0][0] if grid and grid[0] else None
    if title:
        res.fund_name = re.split(r"\s*(분기|운용현황|운용보고)", str(title))[0].strip()

    def row_value(*keys: str) -> int | None:
        for r, c, v in cells(grid):
            if isinstance(v, str) and any(norm(k) in norm(v) for k in keys):
                for x in right_values(grid, r, c):
                    if (n := to_number(x)) is not None:
                        return int(round(n * unit))
        return None

    res.total_asset = row_value("총자산", "AUM")
    res.liabilities = row_value("총부채", "부채총계")
    res.net_asset = row_value("순자산", "NAV")
    if unit != 1:
        res.warnings.append(f"단위 환산 ×{unit:,} (반올림 오차 가능)")


EXTRACTORS = {"A": _extract_a, "B": _extract_b, "C": _extract_c}


def extract_report(data: bytes, filename: str) -> ReportResult:
    res = ReportResult(filename=filename)
    try:
        grid = load_grid(data, filename)
    except Exception as e:  # 손상 파일, 암호 파일 등
        res.warnings.append(f"파일을 열 수 없음: {e}")
        return res
    res.format = detect_format(grid)
    if res.format not in EXTRACTORS:
        res.warnings.append("양식을 인식하지 못함")
        return res
    try:
        EXTRACTORS[res.format](grid, res)
    except Exception as e:
        res.warnings.append(f"추출 중 오류: {e}")
    if res.total_asset is None:
        res.warnings.append("총자산 값을 찾지 못함")
    # 검증식 R11: 자산 − 부채 = 순자산
    if None not in (res.total_asset, res.liabilities, res.net_asset):
        tol = 1 if res.format != "C" else 1_000_000
        if abs(res.total_asset - res.liabilities - res.net_asset) > tol:
            res.warnings.append("대차 불일치(자산−부채≠순자산)")
    return res
