"""추출한 총자산을 당사 수익증권 평가 양식에 채운다 (매칭·제외·누락 처리)."""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date
from difflib import SequenceMatcher

from openpyxl import load_workbook

from .extract import ReportResult, match_key, norm, parse_date

FUZZY_THRESHOLD = 0.85

# 규칙 R9: 양식 하단 주석의 제외 기준 (MMF/ETF, 상장주식, 신종유가증권/초기투자자금)
EXCLUDE_KEYWORDS = ("MMF", "ETF", "코스닥주식", "코스피", "거래소주식", "신종유가증권", "초기투자자금")

STATUS_DONE, STATUS_CHECK, STATUS_FAIL, STATUS_EXCLUDED, STATUS_MISSING = "완료", "확인필요", "실패", "제외", "미수령"


@dataclass
class RowResult:
    row: int
    name: str
    status: str
    total_asset: int | str | None
    source: str
    remark: str


def exclusion_reason(*fields: object) -> str | None:
    text = " ".join(str(f or "") for f in fields).upper()
    for k in EXCLUDE_KEYWORDS:
        if k.upper() in text:
            return k
    if "상장주식" in text.replace("비상장주식", ""):
        return "상장주식"
    return None


def _find_header(ws) -> tuple[int, dict[str, int]]:
    for row in ws.iter_rows(min_row=1, max_row=20):
        names = {norm(c.value): c.column for c in row if c.value is not None}
        if "종목명" in names:
            return row[0].row, names
    raise ValueError("당사 양식에서 '종목명' 헤더를 찾지 못했습니다.")


def _col(headers: dict[str, int], *keys: str) -> int | None:
    for k in keys:
        for h, c in headers.items():
            if norm(k) in h:
                return c
    return None


def load_mapping(data: bytes | None) -> dict[str, dict[str, str]]:
    """매핑표(종목명·표준코드·펀드코드) → {정규화 종목명 or 표준코드: 행}"""
    if not data:
        return {}
    ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = [norm(h) for h in rows[0]]
    out: dict[str, dict[str, str]] = {}
    for r in rows[1:]:
        rec = {h: ("" if v is None else str(v)) for h, v in zip(head, r)}
        for key in (match_key(rec.get("종목명")), rec.get("표준코드", "")):
            if key:
                out[key] = rec
    return out


def _match(name: str, std_code: str, mapping: dict, reports: list[ReportResult]):
    """규칙 R7: 펀드코드 → 정규화 이름 → 유사도 순. (report, 방식) 반환"""
    rec = mapping.get(std_code) or mapping.get(match_key(name)) or {}
    fund_code = rec.get("펀드코드", "")
    names = {match_key(name), match_key(rec.get("종목명"))} - {""}
    if fund_code:
        for rp in reports:
            if rp.fund_code and rp.fund_code == fund_code:
                return rp, "펀드코드"
    for rp in reports:
        if rp.fund_name and match_key(rp.fund_name) in names:
            return rp, "종목명"
    best, score = None, 0.0
    for rp in reports:
        if rp.fund_name:
            s = max(SequenceMatcher(None, match_key(rp.fund_name), n).ratio() for n in names)
            if s > score:
                best, score = rp, s
    if best and score >= FUZZY_THRESHOLD:
        return best, f"유사도 {score:.0%}"
    return None, ""


def fill_form(form_data: bytes, reports: list[ReportResult], mapping_data: bytes | None = None,
              overwrite: bool = False) -> tuple[bytes, list[RowResult], list[ReportResult]]:
    """당사 양식에 총자산을 채운 엑셀(bytes), 행별 결과, 매칭 안 된 보고서를 반환."""
    wb = load_workbook(io.BytesIO(form_data))
    ws = wb.worksheets[0]
    hdr_row, headers = _find_header(ws)
    c_name = _col(headers, "종목명")
    c_code = _col(headers, "표준코드")
    c_date = _col(headers, "기준일")
    c_total = _col(headers, "총자산")
    if c_total is None:
        raise ValueError("당사 양식에서 '총자산' 컬럼을 찾지 못했습니다.")
    c_remark = _col(headers, "비고", "AI처리결과")
    if c_remark is None:  # 비고 칸이 없으면 맨 끝에 추가
        c_remark = ws.max_column + 1
        ws.cell(hdr_row, c_remark, "AI처리결과")
    exclude_cols = [c for c in (_col(headers, "유가증권분류"), _col(headers, "구조화기업성격"),
                                _col(headers, "G/L계정명"), c_name) if c]

    mapping = load_mapping(mapping_data)
    usable = [r for r in reports if r.ok]
    used: set[int] = set()
    results: list[RowResult] = []

    for r in range(hdr_row + 1, ws.max_row + 1):
        name = ws.cell(r, c_name).value
        if not name or (c_date and ws.cell(r, c_date).value is None):
            continue  # 빈 행, 하단 주석 행
        name = str(name)
        std = str(ws.cell(r, c_code).value or "") if c_code else ""
        base: date | None = parse_date(ws.cell(r, c_date).value) if c_date else None
        current = ws.cell(r, c_total).value

        if reason := exclusion_reason(*(ws.cell(r, c).value for c in exclude_cols)):
            value, status, src, remark = "제외", STATUS_EXCLUDED, "", f"제외 대상({reason})"
        else:
            rp, how = _match(name, std, mapping, usable)
            if rp is None:
                value, status, src, remark = None, STATUS_MISSING, "", "보고서 미수령"
            else:
                used.add(id(rp))
                value, src = rp.total_asset, f"{rp.filename} [{how}]"
                notes = list(rp.warnings)
                if how.startswith("유사도"):
                    notes.append(f"이름 유사 매칭({how})")
                if base and rp.base_date and base != rp.base_date:
                    notes.append(f"기준일 불일치(보고서 {rp.base_date})")
                if rp.base_date is None:
                    notes.append("보고서 기준일 확인 불가")
                needs_check = any(k in n for n in notes for k in ("유사", "불일치", "확인 불가", "대차"))
                status = STATUS_CHECK if needs_check else STATUS_DONE
                remark = "; ".join(notes)

        # 규칙 R13: 기존 값은 덮어쓰지 않음
        if current not in (None, "") and not overwrite and current != value:
            remark = f"기존값 유지(추출값 {value})" + (f"; {remark}" if remark else "")
            if status == STATUS_DONE:
                status = STATUS_CHECK
        elif value is not None:
            ws.cell(r, c_total, value)
            if isinstance(value, int):
                ws.cell(r, c_total).number_format = "#,##0"
        ws.cell(r, c_remark, remark or None)
        results.append(RowResult(r, name, status, value, src, remark))

    unmatched = [rp for rp in reports if id(rp) not in used]
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), results, unmatched
